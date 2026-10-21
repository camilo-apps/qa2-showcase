"""
A compact version of the QA2 flow, end to end:

    Traffic Sheet (colours = intent) ─┐
                                      ├─> match by placement ─> rules ─> findings ─> verdict ─> Excel report
    Ad-server export (implemented) ───┘

The production engine reads real Traffic Sheet layouts, talks to the
ad server and runs 31 rules. This demo keeps the same backbone -- the
real `findings`, `verdict`, `colors`, `normalize`, `review` and
`report_name` modules -- with a handful of rules over synthetic data.
"""
from __future__ import annotations

from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from qa2.colors import GREEN, RED, SCOPE_EXCLUDED, WHITE, classify
from qa2.findings import Capability, Domain, EntityType, FindingsBuffer, Status
from qa2.normalize import dims_from_name, norm_dims, normalize_weights
from qa2.report_name import report_filename
from qa2.review import group_review_findings
from qa2.verdict import VERDICT_LABELS, SCOPE_LABELS, headline, overall, verdicts_by_scope


# ───────────────────────────────────────────────────── inputs

@dataclass
class Requested:
    """One Traffic Sheet row, with what its colours say was asked for."""
    placement_id: str
    placement_name: str
    size: str
    creative: str
    start: date | None
    end: date | None
    rotation: object
    verification: str
    intent: str                                  # family of the creative cell
    changed: set[str] = field(default_factory=set)  # columns painted green/red


@dataclass
class Implemented:
    placement_id: str
    placement_name: str
    size: str
    creative: str
    start: date | None
    end: date | None
    rotation: object
    verification: str


def _family(cell) -> str:
    fill = cell.fill
    if fill is None or fill.fill_type != "solid":
        return WHITE
    rgb = str(fill.fgColor.rgb or "")[-6:]
    family, _confidence = classify(rgb)
    return family


def _as_date(value) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    return value if isinstance(value, date) else None


def read_traffic_sheet(path: Path) -> tuple[list[Requested], int]:
    """Rows in scope, and how many were left out (hold / not worked)."""
    sheet = load_workbook(path).active
    header = [c.value for c in sheet[1]]
    col = {name: i for i, name in enumerate(header)}
    rows, excluded = [], 0
    for cells in sheet.iter_rows(min_row=2):
        families = {header[i]: _family(c) for i, c in enumerate(cells)}
        if any(f in SCOPE_EXCLUDED for f in families.values()):
            excluded += 1          # yellow = on hold, grey = not worked: zero findings
            continue
        v = lambda name: cells[col[name]].value
        rows.append(Requested(
            placement_id=str(v("Placement ID")),
            placement_name=v("Placement Name"),
            size=str(v("Size")),
            creative=v("Creative"),
            start=_as_date(v("Start")),
            end=_as_date(v("End")),
            rotation=v("Rotation (%)"),
            verification=v("Verification") or "",
            intent=families["Creative"],
            changed={h for h, f in families.items() if f in (GREEN, RED)},
        ))
    return rows, excluded


def read_export(path: Path) -> list[Implemented]:
    sheet = load_workbook(path).active
    header = [c.value for c in sheet[1]]
    out = []
    for values in sheet.iter_rows(min_row=2, values_only=True):
        r = dict(zip(header, values))
        out.append(Implemented(
            placement_id=str(r["Placement_ID"]),
            placement_name=r["Placement_Name"],
            size=str(r["Placement_Size"]),
            creative=r["Creative_Name"],
            start=_as_date(r["Start_Date"]),
            end=_as_date(r["End_Date"]),
            rotation=r["Rotation"],
            verification=r["Verification_Partner"] or "",
        ))
    return out


# ───────────────────────────────────────────────────── rules

def run_rules(requested: list[Requested], implemented: list[Implemented], campaign: str) -> FindingsBuffer:
    by_placement_ts: dict[str, list[Requested]] = defaultdict(list)
    by_placement_ex: dict[str, list[Implemented]] = defaultdict(list)
    for row in requested:
        by_placement_ts[row.placement_id].append(row)
    for row in implemented:
        by_placement_ex[row.placement_id].append(row)

    numeric_rotation = all(
        isinstance(r.rotation, (int, float)) for r in implemented
    )
    caps_reason = "the rotation weight lives inside the decision tree; no export carries it"
    buffer = FindingsBuffer(campaign=campaign)
    buffer.capabilities.declare(Capability.CREATIVE_EXPORT, True)
    buffer.capabilities.declare(Capability.ROTATION_WEIGHT, numeric_rotation, caps_reason)

    for pid, ts_rows in by_placement_ts.items():
        ex_rows = by_placement_ex.get(pid, [])
        name = ts_rows[0].placement_name
        live = {r.creative: r for r in ex_rows}
        ctx = dict(placement_id=pid, placement_name=name)

        if not ex_rows:
            buffer.not_verified("PLC-001", Domain.IDENTITY,
                                "Placement is not in the export, nothing could be compared.",
                                reason="placement missing from the export", **ctx)
            continue

        for row in ts_rows:
            cre = dict(ctx, creative_name=row.creative, entity_type=EntityType.CREATIVE)

            # CRE-001 / CRE-002 -- green must be live, red must be gone.
            if row.intent == GREEN:
                if row.creative in live:
                    buffer.pass_("CRE-001", Domain.CREATIVE, "Requested creative is assigned.", **cre)
                else:
                    others = ", ".join(sorted(set(live) - {r.creative for r in ts_rows})) or "none"
                    buffer.fail("CRE-001", Domain.CREATIVE,
                                "Requested creative is not assigned on this placement.",
                                expected=row.creative, actual=others, **cre)
            elif row.intent == RED:
                if row.creative in live:
                    buffer.fail("CRE-002", Domain.CREATIVE,
                                "Creative marked for removal is still assigned.",
                                expected="(unassigned)", actual=row.creative, **cre)
                else:
                    buffer.pass_("CRE-002", Domain.CREATIVE, "Removed creative is unassigned.", **cre)
                continue

            # DAT-001 -- an end date only one side has is a question, not an error.
            got = live.get(row.creative)
            if got and "End" in row.changed:
                if got.end == row.end:
                    buffer.pass_("DAT-001", Domain.DATES, "End date matches.", **cre)
                elif got.end is None or row.end is None:
                    buffer.review("DAT-001", Domain.DATES,
                                  f"{row.creative} has no end date in the ad server, "
                                  f"and the Traffic Sheet closes it on {row.end}.",
                                  expected=str(row.end), actual=str(got.end or "(none)"), **cre)
                else:
                    buffer.fail("DAT-001", Domain.DATES, "End date differs.",
                                expected=str(row.end), actual=str(got.end), **cre)

        # DIM-001 -- every live creative must fit the placement it serves on.
        for ex in ex_rows:
            dims = dims_from_name(ex.creative)
            if dims and is_display(ex.size) and norm_dims(dims) != norm_dims(ex.size):
                buffer.fail("DIM-001", Domain.DIMENSIONS,
                            f"A {dims} creative is serving on a {ex.size} placement.",
                            expected=ex.size, actual=f"{ex.creative} ({dims})",
                            creative_name=ex.creative, **ctx)

        # ROT-001 -- requested weights vs implemented weights.
        if any("Rotation (%)" in r.changed for r in ts_rows):
            if not buffer.capabilities.available(Capability.ROTATION_WEIGHT):
                # What could not be checked never passes green.
                buffer.not_verified("ROT-001", Domain.ROTATION,
                                    "Requested rotation weights were not compared.",
                                    reason=buffer.capabilities.reason(Capability.ROTATION_WEIGHT),
                                    expected=" / ".join(normalize_weights([r.rotation for r in ts_rows])),
                                    entity_type=EntityType.ROTATION, **ctx)
            else:
                _compare_rotation(buffer, ts_rows, live, ctx)

        # TAG-001 -- verification partner, counted in the tags & pixels scope.
        wanted = {r.verification for r in ts_rows if r.intent != RED}
        found = {r.verification for r in ex_rows}
        if wanted <= found:
            buffer.pass_("TAG-001", Domain.TAG, "Verification partner is in place.",
                         expected=", ".join(sorted(wanted)), actual=", ".join(sorted(found)), **ctx)
        else:
            buffer.fail("TAG-001", Domain.TAG, "Verification partner is missing.",
                        expected=", ".join(sorted(wanted)), actual=", ".join(sorted(found)) or "(none)", **ctx)

    return buffer


def _compare_rotation(buffer: FindingsBuffer, ts_rows, live, ctx) -> None:
    # Each side stores weights in its own scale (0.5 vs 50 vs shares);
    # normalize_weights turns both into percentages that add up to 100.
    asked = normalize_weights([r.rotation for r in ts_rows])
    done = normalize_weights([live[r.creative].rotation if r.creative in live else "" for r in ts_rows])
    status = Status.PASS if asked == done else Status.FAIL
    buffer.emit("ROT-001", Domain.ROTATION, status,
                "Rotation weights match." if status is Status.PASS else "Rotation weights differ.",
                expected=" / ".join(asked), actual=" / ".join(done),
                entity_type=EntityType.ROTATION,
                requires=[Capability.ROTATION_WEIGHT], **ctx)


def is_display(size: str) -> bool:
    return bool(dims_from_name(size)) and size not in ("16x9", "9x16", "1x1")


# ───────────────────────────────────────────────────── sign-off + report

@dataclass
class SignOff:
    who: str
    why: str
    when: str


def write_report(findings, campaign: str, folder: Path, signoffs: dict[str, SignOff] | None = None,
                 excluded_rows: int = 0, when: datetime | None = None) -> Path:
    signoffs = signoffs or {}
    folder.mkdir(parents=True, exist_ok=True)
    path = folder / report_filename("xlsx", campaign, when)
    colours = {"PASS": "E7F7F1", "FAIL": "FDECEC", "REVIEW": "FDF3E3", "NOT_VERIFIED": "EEF1F6"}

    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    verdict = overall(findings)
    summary.append(["Campaign", campaign])
    summary.append(["Overall result", VERDICT_LABELS[verdict]])
    for scope, value in sorted(verdicts_by_scope(findings).items()):
        summary.append([SCOPE_LABELS[scope], VERDICT_LABELS[value]])
    if headline(findings):
        summary.append(["What to do", headline(findings)])
    summary.append(["Rows out of scope (hold / not worked)", excluded_rows])
    for row in summary.iter_rows():
        row[0].font = Font(bold=True)
    summary.column_dimensions["A"].width = 38
    summary.column_dimensions["B"].width = 70

    sheet = wb.create_sheet("Findings")
    sheet.append(["Status", "Rule", "Placement ID", "Placement", "Creative",
                  "Message", "Expected", "Actual", "Signed by", "Why", "When"])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for f in findings:
        sign = signoffs.get(f.finding_id)
        sheet.append([f.status.value, f.rule_id, f.placement_id, f.placement_name, f.creative_name,
                      f.message, f.expected, f.actual,
                      sign.who if sign else "", sign.why if sign else "", sign.when if sign else ""])
        fill = colours.get(f.status.value)
        if fill:
            sheet.cell(row=sheet.max_row, column=1).fill = PatternFill("solid", fgColor=fill)
    for col, width in zip("ABCDEFGHIJK", (14, 9, 12, 32, 22, 60, 22, 26, 16, 34, 18)):
        sheet.column_dimensions[col].width = width
    sheet.freeze_panes = "A2"
    wb.save(path)
    return path


def waiting_for_signature(findings) -> dict:
    return group_review_findings(f for f in findings if f.status is Status.REVIEW)
