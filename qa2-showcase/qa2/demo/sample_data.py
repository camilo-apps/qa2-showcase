"""
Synthetic inputs for the demo: a Traffic Sheet and an ad-server export.

Everything here is invented -- the advertiser, the placement IDs, the
creatives and the dates. The shapes mirror the real problem: the
Traffic Sheet says *what was requested* through cell colours, and the
export says *what was actually implemented*.

The data carries the cases the demo is meant to show:

  * a creative swap done right                    -> PASS
  * two rotations crossed between sizes           -> FAIL (x2)
  * a creative removed in the sheet, still live   -> FAIL
  * an end date only one side has                 -> REVIEW (needs a human)
  * a rotation stored in a decision tree          -> NOT VERIFIED, never green
  * rows on hold (yellow) and not worked (grey)   -> out of scope, no findings
"""
from __future__ import annotations

from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

GREEN = PatternFill("solid", fgColor="C6EFCE")
RED = PatternFill("solid", fgColor="FFC7CE")
YELLOW = PatternFill("solid", fgColor="FFEB9C")
GREY = PatternFill("solid", fgColor="D9D9D9")

CAMPAIGN = "Northwind Coffee Holiday 2026"

TS_HEADER = [
    "Placement ID", "Placement Name", "Size", "Creative",
    "Start", "End", "Rotation (%)", "Verification",
]

# (placement id, name, size, creative, start, end, rotation, verification, fill per column)
# fill per column: {column header: fill}
TS_ROWS = [
    ("40021811", "NWC_Holiday_Display_728x90", "728x90", "Mocha_728x90",
     date(2026, 11, 15), date(2026, 12, 28), 1.0, "DV", {"Creative": GREEN}),
    ("40021877", "NWC_Holiday_Display_300x600", "300x600", "Latte_300x600",
     date(2026, 11, 15), date(2026, 12, 28), 0.5, "DV", {"Creative": GREEN}),
    ("40021877", "NWC_Holiday_Display_300x600", "300x600", "Mocha_300x600",
     date(2026, 11, 15), date(2026, 12, 28), 0.5, "DV", {}),
    ("40021903", "NWC_Holiday_Display_300x250", "300x250", "Latte_300x250",
     date(2026, 11, 15), date(2026, 12, 28), 0.5, "DV", {"Creative": GREEN}),
    ("40021903", "NWC_Holiday_Display_300x250", "300x250", "Mocha_300x250",
     date(2026, 11, 15), date(2026, 12, 28), 0.5, "DV", {}),
    ("40021955", "NWC_Holiday_OLV_16x9", "16x9", "Holiday_Spot_15s",
     date(2026, 11, 15), date(2026, 12, 28), 1.0, "DV", {"End": GREEN}),
    ("40021960", "NWC_Holiday_Display_160x600", "160x600", "Summer_160x600",
     date(2026, 6, 1), date(2026, 8, 31), "", "DV", {"Creative": RED}),
    ("40021960", "NWC_Holiday_Display_160x600", "160x600", "Mocha_160x600",
     date(2026, 11, 15), date(2026, 12, 28), 1.0, "DV", {"Creative": GREEN}),
    ("40021836", "NWC_Holiday_CTV_30s", "16x9", "Holiday_Spot_30s",
     date(2026, 11, 15), date(2026, 12, 28), 0.6, "DV", {"Rotation (%)": GREEN}),
    ("40021836", "NWC_Holiday_CTV_30s", "16x9", "Holiday_Spot_30s_B",
     date(2026, 11, 15), date(2026, 12, 28), 0.4, "DV", {"Rotation (%)": GREEN}),
    ("40021990", "NWC_Holiday_Display_320x50", "320x50", "Latte_320x50",
     date(2026, 11, 20), date(2026, 12, 28), 1.0, "DV", {"*": YELLOW}),
    ("40021700", "NWC_Always_On_Display_300x250", "300x250", "Brand_300x250",
     date(2026, 1, 1), date(2026, 12, 31), 1.0, "DV", {"*": GREY}),
]

EXPORT_HEADER = [
    "Placement_ID", "Placement_Name", "Placement_Size", "Creative_Name",
    "Start_Date", "End_Date", "Rotation", "Verification_Partner",
]

# What the ad server says was implemented.
EXPORT_ROWS = [
    ("40021811", "NWC_Holiday_Display_728x90", "728x90", "Mocha_728x90",
     date(2026, 11, 15), date(2026, 12, 28), 100, "DV"),
    # The two rotations below are crossed: each placement carries the
    # other size's Latte creative.
    ("40021877", "NWC_Holiday_Display_300x600", "300x600", "Latte_300x250",
     date(2026, 11, 15), date(2026, 12, 28), 50, "DV"),
    ("40021877", "NWC_Holiday_Display_300x600", "300x600", "Mocha_300x600",
     date(2026, 11, 15), date(2026, 12, 28), 50, "DV"),
    ("40021903", "NWC_Holiday_Display_300x250", "300x250", "Latte_300x600",
     date(2026, 11, 15), date(2026, 12, 28), 50, "DV"),
    ("40021903", "NWC_Holiday_Display_300x250", "300x250", "Mocha_300x250",
     date(2026, 11, 15), date(2026, 12, 28), 50, "DV"),
    # No end date on the ad server side.
    ("40021955", "NWC_Holiday_OLV_16x9", "16x9", "Holiday_Spot_15s",
     date(2026, 11, 15), None, 100, "DV"),
    # The red creative was never unassigned.
    ("40021960", "NWC_Holiday_Display_160x600", "160x600", "Summer_160x600",
     date(2026, 6, 1), date(2026, 8, 31), 50, "DV"),
    ("40021960", "NWC_Holiday_Display_160x600", "160x600", "Mocha_160x600",
     date(2026, 11, 15), date(2026, 12, 28), 50, "DV"),
    # Rotation lives inside a decision tree: no export carries the weight.
    ("40021836", "NWC_Holiday_CTV_30s", "16x9", "Holiday_Spot_30s",
     date(2026, 11, 15), date(2026, 12, 28), "Decision Tree", "DV"),
    ("40021836", "NWC_Holiday_CTV_30s", "16x9", "Holiday_Spot_30s_B",
     date(2026, 11, 15), date(2026, 12, 28), "Decision Tree", "DV"),
]


def write_samples(folder: Path) -> tuple[Path, Path]:
    folder.mkdir(parents=True, exist_ok=True)

    ts = Workbook()
    sheet = ts.active
    sheet.title = "Traffic"
    sheet.append(TS_HEADER)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for *values, fills in TS_ROWS:
        sheet.append(values)
        row = sheet.max_row
        for col, header in enumerate(TS_HEADER, start=1):
            fill = fills.get(header) or fills.get("*")
            if fill:
                sheet.cell(row=row, column=col).fill = fill
    ts_path = folder / "traffic_sheet.xlsx"
    ts.save(ts_path)

    ex = Workbook()
    sheet = ex.active
    sheet.title = "Placement-Creative"
    sheet.append(EXPORT_HEADER)
    for values in EXPORT_ROWS:
        sheet.append(list(values))
    export_path = folder / "placement_creative_export.xlsx"
    ex.save(export_path)

    return ts_path, export_path
