"""
The demo, end to end, on the synthetic data in qa2.demo.sample_data.

Each test is one of the cases the demo exists to show.
"""
import sys
from datetime import datetime
from pathlib import Path

import pytest
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from qa2.demo.pipeline import (  # noqa: E402
    SignOff, read_export, read_traffic_sheet, run_rules, waiting_for_signature, write_report,
)
from qa2.demo.sample_data import CAMPAIGN, write_samples  # noqa: E402
from qa2.findings import Capability, CapabilityRegistry, Domain, FindingsBuffer, Status, Confidence  # noqa: E402
from qa2.verdict import IMPLEMENTATION, TAGS, overall, verdicts_by_scope  # noqa: E402


@pytest.fixture(scope="module")
def run(tmp_path_factory):
    folder = tmp_path_factory.mktemp("demo")
    ts, export = write_samples(folder)
    requested, excluded = read_traffic_sheet(ts)
    findings = run_rules(requested, read_export(export), CAMPAIGN).findings
    return folder, requested, excluded, findings


def _by(findings, rule, placement):
    return [f for f in findings if f.rule_id == rule and f.placement_id == placement]


def test_crossed_rotations_are_caught_on_both_placements(run):
    *_, findings = run
    for placement in ("40021877", "40021903"):
        dims = _by(findings, "DIM-001", placement)
        assert [f.status for f in dims] == [Status.FAIL]
        assert _by(findings, "CRE-001", placement)[0].status is Status.FAIL


def test_a_red_creative_still_live_fails(run):
    *_, findings = run
    assert _by(findings, "CRE-002", "40021960")[0].status is Status.FAIL


def test_a_one_sided_end_date_goes_to_a_human(run):
    *_, findings = run
    [date_check] = _by(findings, "DAT-001", "40021955")
    assert date_check.status is Status.REVIEW
    assert "no end date in the ad server" in date_check.message


def test_what_could_not_be_checked_never_passes_green(run):
    *_, findings = run
    [rotation] = _by(findings, "ROT-001", "40021836")
    assert rotation.status is Status.NOT_VERIFIED
    assert "decision tree" in rotation.reason


def test_rows_on_hold_or_not_worked_produce_no_findings(run):
    _, requested, excluded, findings = run
    assert excluded == 2
    placements = {f.placement_id for f in findings}
    assert "40021990" not in placements   # yellow: on hold
    assert "40021700" not in placements   # grey: not worked


def test_implementation_and_tags_get_their_own_verdict(run):
    *_, findings = run
    verdicts = verdicts_by_scope(findings)
    assert verdicts[IMPLEMENTATION] == "FAILED"
    assert verdicts[TAGS] == "PASSED"
    assert overall(findings) == "FAILED"


def test_the_report_carries_the_signature(run, tmp_path):
    *_, findings = run
    [(key, items)] = waiting_for_signature(findings).items()
    signoffs = {f.finding_id: SignOff("Laura Gómez", "Client extended the flight", "2026-10-21 10:42")
                for f in items}
    path = write_report(findings, CAMPAIGN, tmp_path, signoffs, when=datetime(2026, 10, 21))
    assert path.name == "2026-10-21_QA_Report_Northwind_Coffee_Holiday_2026.xlsx"
    rows = list(load_workbook(path)["Findings"].iter_rows(values_only=True))
    signed = [r for r in rows if r[0] == "REVIEW"]
    assert signed and signed[0][8] == "Laura Gómez"


# ── the two gates in FindingsBuffer, in isolation ──────────────────────

def test_a_missing_input_turns_a_fail_into_not_verified():
    caps = CapabilityRegistry()
    caps.declare(Capability.ROTATION_WEIGHT, False, "no export carries the weight")
    buffer = FindingsBuffer(caps)
    f = buffer.fail("ROT-001", Domain.ROTATION, "Rotation weights differ.",
                    requires=[Capability.ROTATION_WEIGHT])
    assert f.status is Status.NOT_VERIFIED
    assert f.degraded_from is not None


def test_a_low_confidence_fail_goes_to_review():
    buffer = FindingsBuffer()
    f = buffer.fail("CRE-001", Domain.CREATIVE, "Creative not found.", confidence=Confidence.MEDIUM)
    assert f.status is Status.REVIEW
